"""검수확인서(제품 검수 확인서) 엑셀 자동 생성.

template_files/inspection_template.xlsx (회사 제공 양식)을 원본으로, 견적서/
지출결의서에서 얻은 값으로 "내용작성" 자리를 채운다. 양식에 박혀 있는 로고/
서명 등 이미지와 서식은 그대로 보존한다(openpyxl load->edit->save).

양식 좌표(고정):
  F4 = 회사명,  J4 = 주소,  F5 = 업종,  J5 = 기타
  제품 내역: 7~13행 (E=품목, F=규격, G:H=수량, I=공급가, J=비고)
  F14/F15 = 검수 항목 결과("이상 없음" 기본)
  B17:J39 = 제품 사진(비워둠 - 수기 첨부)
  F40 = 검수 부서(그린연구소), J40 = 검수자 성명, F42 = 검수일

품목이 7개를 넘으면 제품 내역 표 아래(14행 앞)에 필요한 만큼 행을 삽입한다.
"""

import copy
import datetime
import io
import os

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from paths import bundle_dir

TEMPLATE_PATH = os.path.join(bundle_dir(), "template_files", "inspection_template.xlsx")

# 제품 내역 표 (양식 고정 좌표)
ITEM_FIRST_ROW = 7
ITEM_LAST_ROW = 13          # 기본 제공 칸의 마지막 행
BASE_ITEM_SLOTS = ITEM_LAST_ROW - ITEM_FIRST_ROW + 1  # 7
COL_NAME = "E"
COL_SPEC = "F"
COL_QTY = "G"               # G:H 병합
COL_SUPPLY = "I"
COL_NOTE = "J"

# 납품자 정보
CELL_COMPANY = "F4"
CELL_ADDRESS = "J4"
CELL_INDUSTRY = "F5"
CELL_ETC = "J5"

# 검수 항목 결과
CELL_SPEC_CHECK = "F14"
CELL_QTY_CHECK = "F15"
DEFAULT_CHECK_RESULT = "이상 없음"

# 검수자 확인
CELL_DEPARTMENT = "F40"
CELL_INSPECTOR = "J40"
CELL_SIGN = "F41"
CELL_INSPECT_DATE = "F42"

DEFAULT_DEPARTMENT = "그린연구소"

_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True, shrink_to_fit=True)


def _to_date(value):
    if not value:
        return None
    if isinstance(value, (datetime.date, datetime.datetime)):
        return value
    return datetime.datetime.strptime(str(value), "%Y-%m-%d").date()


def item_supply_amount(item):
    """품목 한 줄의 공급가(금액)를 계산한다. 지출결의서 계산 규칙과 동일.

    - 단가가 있으면: 중량(있고 0이 아니면) 또는 수량 x 단가
    - 단가가 없으면: item['supply'] (인쇄된 금액) 그대로
    """
    price = item.get("price")
    if price:
        weight = item.get("weight")
        qty = item.get("qty")
        if weight:
            return round(price * weight)
        if qty:
            return round(price * qty)
        return round(price)
    supply = item.get("supply")
    if supply:
        return round(supply)
    return 0


def _shift_images_down(ws, from_row, n_rows):
    """from_row(1-indexed) 이상에 앵커된 이미지를 n_rows 만큼 아래로 옮긴다.

    행을 삽입하면 openpyxl 은 이미지 앵커를 자동으로 밀지 않으므로 수동 보정.
    """
    for img in getattr(ws, "_images", []):
        try:
            anchor = img.anchor._from  # 0-indexed row
        except AttributeError:
            continue
        if anchor.row >= from_row - 1:
            anchor.row += n_rows
            to = getattr(img.anchor, "to", None)
            if to is not None:
                to.row += n_rows


def _insert_item_rows(ws, extra):
    """제품 내역 표에 extra 개의 행을 추가한다(ITEM_LAST_ROW 다음에 삽입).

    openpyxl insert_rows 는 셀 값/서식만 밀고 병합/이미지는 밀지 않으므로,
    삽입 지점(14행) 이상의 병합 범위와 이미지 앵커를 수동으로 함께 내린다.
    """
    insert_at = ITEM_LAST_ROW + 1  # 14

    # 1) insert_at 이상에 걸친 병합을 기록 후 해제
    to_shift = [(m.min_row, m.min_col, m.max_row, m.max_col)
                for m in list(ws.merged_cells.ranges) if m.min_row >= insert_at]
    for min_r, min_c, max_r, max_c in to_shift:
        ws.unmerge_cells(start_row=min_r, start_column=min_c, end_row=max_r, end_column=max_c)

    # 2) 이미지 앵커 이동
    _shift_images_down(ws, insert_at, extra)

    # 3) 행 삽입(값/스타일 이동)
    ws.insert_rows(insert_at, extra)

    # 4) 병합 범위 재적용(+extra)
    for min_r, min_c, max_r, max_c in to_shift:
        ws.merge_cells(start_row=min_r + extra, start_column=min_c,
                       end_row=max_r + extra, end_column=max_c)

    # 5) 새 품목 행 서식/높이 복제 + 수량 G:H 병합
    src = ITEM_LAST_ROW  # 마지막 기본 칸(13)의 서식을 복제
    for off in range(extra):
        dst = insert_at + off
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height
        for col in range(2, 11):  # B~J
            s = ws.cell(row=src, column=col)
            d = ws.cell(row=dst, column=col)
            if s.has_style:
                d._style = copy.copy(s._style)
        try:
            ws.merge_cells(f"G{dst}:H{dst}")
        except Exception:
            pass


def build_inspection_report(data):
    """검수확인서 워크북을 채워 BytesIO 로 반환한다.

    data:
      company, address, industry, etc     : 납품자 정보(문자열, 없으면 빈칸)
      items          : [{name, spec?, qty?, price?, weight?, supply?, note?}, ...]
      inspector      : 검수자 성명
      inspect_date   : 'YYYY-MM-DD' 또는 date
      department     : 검수 부서(기본 '그린연구소')
      spec_check     : 규격 확인 결과(기본 '이상 없음')
      qty_check      : 수량 확인 결과(기본 '이상 없음')
    """
    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb.worksheets[0]

    items = data.get("items") or []
    if not items:
        raise ValueError("검수확인서에 넣을 품목이 최소 1개 이상 필요합니다.")

    # 품목이 기본 칸보다 많으면 행 삽입. 그만큼 14행 이하의 고정 좌표가 밀린다.
    extra = max(0, len(items) - BASE_ITEM_SLOTS)
    if extra:
        _insert_item_rows(ws, extra)

    def shift(cell):
        """14행 이상 좌표를 삽입한 행 수(extra)만큼 내린 주소로 바꾼다."""
        col = "".join(ch for ch in cell if ch.isalpha())
        row = int("".join(ch for ch in cell if ch.isdigit()))
        if extra and row >= ITEM_LAST_ROW + 1:
            row += extra
        return f"{col}{row}"

    # 납품자 정보
    ws[CELL_COMPANY] = data.get("company", "") or ""
    ws[CELL_ADDRESS] = data.get("address", "") or ""
    ws[CELL_INDUSTRY] = data.get("industry", "") or ""
    ws[CELL_ETC] = data.get("etc", "") or ""

    # 제품 내역
    row = ITEM_FIRST_ROW
    for item in items:
        ws[f"{COL_NAME}{row}"] = item.get("name", "")
        spec = item.get("spec")
        ws[f"{COL_SPEC}{row}"] = spec if spec and spec != "-" else ("" if not spec else spec)
        qty = item.get("qty")
        qcell = ws[f"{COL_QTY}{row}"]
        qcell.value = qty if qty is not None else ""
        if isinstance(qty, float) and qty.is_integer():
            qcell.value = int(qty)
        supply = item_supply_amount(item)
        scell = ws[f"{COL_SUPPLY}{row}"]
        scell.value = supply
        scell.number_format = "#,##0"
        note = item.get("note")
        if note:
            ws[f"{COL_NOTE}{row}"] = note
        row += 1

    # 빈 칸(품목이 기본 슬롯보다 적을 때 남는 7~13행 자리)의 "내용작성" 잔여물 제거
    last_item_row = ITEM_FIRST_ROW + len(items) - 1
    max_slot_row = max(ITEM_LAST_ROW, last_item_row)
    for r in range(last_item_row + 1, max_slot_row + 1):
        for col in (COL_NAME, COL_SPEC, COL_QTY, COL_SUPPLY, COL_NOTE):
            ws[f"{col}{r}"] = None

    # 검수 항목 결과
    ws[shift(CELL_SPEC_CHECK)] = data.get("spec_check", DEFAULT_CHECK_RESULT)
    ws[shift(CELL_QTY_CHECK)] = data.get("qty_check", DEFAULT_CHECK_RESULT)

    # 검수자 확인
    ws[shift(CELL_DEPARTMENT)] = data.get("department") or DEFAULT_DEPARTMENT
    ws[shift(CELL_INSPECTOR)] = data.get("inspector", "") or ""
    # 서명 칸의 "내용작성" 안내문구는 지우고 수기 서명 자리로 비워둔다
    ws[shift(CELL_SIGN)] = None
    inspect_date = _to_date(data.get("inspect_date"))
    ws[shift(CELL_INSPECT_DATE)] = inspect_date.strftime("%Y-%m-%d") if inspect_date else ""

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def suggest_filename(company, category=None):
    company = (company or "업체명미상").replace(" ", "")
    if category:
        return f"검수확인서_{company}_{category}.xlsx"
    return f"검수확인서_{company}.xlsx"
