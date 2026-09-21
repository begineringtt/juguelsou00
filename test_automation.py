"""신규 자동화 모듈 테스트.

실제 샘플 견적서(setting_03)가 있으면 파서 정확도까지 검증하고, 없으면
합성 데이터로 생성/폴더/체크리스트 로직을 검증한다(샘플 없어도 CI 통과).
"""

import datetime
import io
import os
import unittest

import openpyxl

import attachment_finder
import checklist_updater
import folder_router
import inspection_generator
import report_writer


class FolderRouterTest(unittest.TestCase):
    def test_중동_maps_to_IR(self):
        self.assertEqual(folder_router.folder_for_category("중동"), "IR")
        self.assertEqual(folder_router.folder_for_category("중동(IR)"), "IR")

    def test_고효율_maps_to_광원폴더(self):
        self.assertEqual(folder_router.folder_for_category("고효율"), "고효율 광원")

    def test_date_pattern_new_folder(self):
        sib = ["2026-09-20 유진철강", "2026.09.01 코리아농업개발"]
        name, is_new, pat = folder_router.propose_company_folder(
            sib, "무성전자", today=datetime.date(2026, 9, 21))
        self.assertTrue(is_new)
        self.assertIn("무성전자", name)
        self.assertTrue(name.startswith("2026"))

    def test_cha_pattern_next_number(self):
        sib = ["10차 신안그린테크", "7차 다임", "9차 한수과학"]
        name, is_new, pat = folder_router.propose_company_folder(sib, "가나테크")
        self.assertTrue(is_new)
        self.assertTrue(name.startswith("11차"))

    def test_existing_folder_matched(self):
        sib = ["10차 신안그린테크"]
        name, is_new, pat = folder_router.propose_company_folder(sib, "신안그린테크")
        self.assertFalse(is_new)
        self.assertEqual(name, "10차 신안그린테크")


class InspectionGeneratorTest(unittest.TestCase):
    def _items(self, n):
        return [{"name": f"품목{i}", "spec": f"S{i}", "qty": i, "price": 1000 * i}
                for i in range(1, n + 1)]

    def test_supply_weight_based(self):
        it = {"name": "각파이프", "qty": 90, "weight": 4849.2, "price": 1040}
        self.assertEqual(inspection_generator.item_supply_amount(it), 5043168)

    def test_supply_qty_based(self):
        it = {"name": "밸브", "qty": 3, "price": 1000}
        self.assertEqual(inspection_generator.item_supply_amount(it), 3000)

    def test_build_basic(self):
        buf = inspection_generator.build_inspection_report({
            "company": "테스트㈜", "items": self._items(2),
            "inspector": "홍길동", "inspect_date": "2026-09-21"})
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        ws = wb.active
        self.assertEqual(ws["F4"].value, "테스트㈜")
        self.assertEqual(ws["E7"].value, "품목1")
        self.assertEqual(ws["F14"].value, "이상 없음")
        self.assertEqual(ws["J40"].value, "홍길동")
        self.assertEqual(len(ws._images), 6)  # 양식 이미지 보존

    def test_overflow_rows_shift_footer(self):
        buf = inspection_generator.build_inspection_report({
            "company": "테스트㈜", "items": self._items(9),
            "inspector": "홍길동", "inspect_date": "2026-09-21"})
        wb = openpyxl.load_workbook(io.BytesIO(buf.getvalue()))
        ws = wb.active
        # 9품목 -> extra=2 -> 검수자 확인 라벨이 42행으로 이동
        self.assertEqual(ws["E15"].value, "품목9")
        self.assertEqual(ws["B42"].value, "검수자 확인")
        self.assertEqual(ws["J42"].value, "홍길동")


class ChecklistUpdaterTest(unittest.TestCase):
    def _make_checklist(self, path):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        # 헤더: 업체 열 (row1, col D..)
        ws["D1"] = "유진철강"
        ws["E1"] = "그린공조"
        # 카테고리 블록: 고효율 (row5..11)
        docs = ["1. 견적서", "2. 사업자등록증", "3. 통장사본", "4. 지출결의서",
                "5. 전자세금계산서", "6. 거래명세서", "7. 검수확인서"]
        ws["A5"] = "고효율"
        for i, d in enumerate(docs):
            ws.cell(row=5 + i, column=2, value=d)
        wb.save(path)

    def test_mark_cells(self):
        src = "/tmp/_ck_src.xlsx"
        out = "/tmp/_ck_out.xlsx"
        self._make_checklist(src)
        res = checklist_updater.update_checklist(
            src, "고효율", "유진철강산업㈜",
            ["견적서", "지출결의서", "검수확인서"], out_path=out)
        self.assertEqual(res["misses"], [])
        cells = {m["cell"] for m in res["marked"]}
        self.assertIn("D5", cells)   # 견적서
        self.assertIn("D8", cells)   # 지출결의서
        self.assertIn("D11", cells)  # 검수확인서
        wb = openpyxl.load_workbook(out)
        self.assertEqual(wb["Sheet1"]["D5"].value, "O")


class ReportWriterTest(unittest.TestCase):
    def test_append_creates_and_grows(self):
        path = "/tmp/_report.xlsx"
        if os.path.exists(path):
            os.remove(path)
        report_writer.append_entry(path, {
            "category": "고효율", "company": "A㈜", "folder": "고효율 광원/x",
            "generated": ["지출결의서"], "attachments": ["견적서"],
            "total_amount": 1000, "item_count": 1, "source": "pdf-table", "note": ""})
        report_writer.append_entry(path, {
            "category": "중동", "company": "B㈜", "folder": "IR/y",
            "generated": ["검수확인서"], "attachments": [],
            "total_amount": 2000, "item_count": 2, "source": "xlsx", "note": "재확인"})
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        self.assertEqual(ws.max_row, 3)  # 헤더 + 2행
        self.assertEqual(ws["C2"].value, "A㈜")
        self.assertEqual(ws["C3"].value, "B㈜")


class AttachmentFinderTest(unittest.TestCase):
    def _make_repo(self, root):
        # 다른 과제 폴더에 유진철강의 통장/사업자가 있는 상황
        d = os.path.join(root, "AI(1y)그린CS", "3차 유진철강")
        os.makedirs(d)
        open(os.path.join(d, "유진철강산업통장사본1.pdf"), "w").close()
        open(os.path.join(d, "20230616 유진철강산업 사업자.pdf"), "w").close()
        d2 = os.path.join(root, "고온성", "그린공조시스템")
        os.makedirs(d2)
        open(os.path.join(d2, "6. 통장사본_그린공조시스템.jpg"), "w").close()

    def test_index_and_find(self):
        import tempfile
        root = tempfile.mkdtemp(prefix="repo_")
        self._make_repo(root)
        idx = attachment_finder.build_index(root)
        docs, disp = attachment_finder.find_documents(idx, "유진철강산업㈜")
        self.assertIsNotNone(docs["통장사본"])
        self.assertIn("통장", os.path.basename(docs["통장사본"]))
        self.assertIsNotNone(docs["사업자등록증"])
        # 폴더명 변형 매칭
        docs2, _ = attachment_finder.find_documents(idx, "그린공조")
        self.assertIsNotNone(docs2["통장사본"])
        # 없는 업체
        docs3, disp3 = attachment_finder.find_documents(idx, "존재하지않는회사")
        self.assertIsNone(docs3["통장사본"])
        self.assertIsNone(disp3)


# ---- 실제 샘플이 있을 때만 도는 파서 정확도 테스트 ----
SAMPLE = "/mnt/user-data/uploads/setting_03/고효율 광원/2026-09-20 유진철강/견적서.pdf"


@unittest.skipUnless(os.path.isfile(SAMPLE), "실제 견적서 샘플 없음")
class QuoteReaderSampleTest(unittest.TestCase):
    def test_유진철강_coords(self):
        import quote_reader
        res = quote_reader.read_quote(path=SAMPLE)
        self.assertEqual(len(res["items"]), 1)
        it = res["items"][0]
        self.assertEqual(it["name"], "각파이프")
        self.assertEqual(it["qty"], 90.0)
        self.assertEqual(it["price"], 1040.0)
        self.assertAlmostEqual(it["weight"], 4849.2)


if __name__ == "__main__":
    unittest.main(verbosity=2)
