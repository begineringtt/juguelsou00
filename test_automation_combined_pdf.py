import os
import shutil
import tempfile

import openpyxl
from PIL import Image

import automation


def _make_quote_xlsx(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["품명", "수량", "단가"])
    ws.append(["테스트자재", 2, 1000])
    wb.save(path)


def _make_image(path):
    Image.new("RGB", (200, 100), color="white").save(path)


def test_run_places_combined_pdf_in_target_folder_and_reports_it():
    tmp_dir = tempfile.mkdtemp()
    try:
        setting03_root = os.path.join(tmp_dir, "setting_03")
        os.makedirs(os.path.join(setting03_root, "고온성"))

        quote_path = os.path.join(tmp_dir, "견적서.xlsx")
        biz_path = os.path.join(tmp_dir, "사업자등록증.jpg")
        bank_path = os.path.join(tmp_dir, "통장사본.jpg")
        _make_quote_xlsx(quote_path)
        _make_image(biz_path)
        _make_image(bank_path)

        manifest = automation.run(
            quote_path, "고온성", "테스트업체", setting03_root,
            attachments={"사업자등록증": biz_path, "통장사본": bank_path},
            auto_find_attachments=False, update_checklist=False, write_report=True,
        )

        assert manifest["combined_pdf"]
        combined_name = os.path.basename(manifest["combined_pdf"])
        target_dir = manifest["target_dir"]
        assert os.path.isfile(os.path.join(target_dir, combined_name))
        assert any(p.endswith(combined_name) for p in manifest["placed_files"])

        report_path = manifest["report_path"]
        wb = openpyxl.load_workbook(report_path)
        ws = wb.active
        headers = [c.value for c in ws[1]]
        assert "통합PDF" in headers
        col = headers.index("통합PDF") + 1
        assert ws.cell(row=2, column=col).value == "O"
        print("OK: test_run_places_combined_pdf_in_target_folder_and_reports_it")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)


if __name__ == "__main__":
    test_run_places_combined_pdf_in_target_folder_and_reports_it()
    print("ALL PASSED")
