"""연구비 파일 업로드 체크용 엑셀 갱신.

체크리스트 구조(파악 결과):
  - 카테고리 블록이 세로로 이어짐: 각 블록 = 7행
      (row0) 1. 견적서 / 2. 사업자등록증 / 3. 통장사본 / 4. 지출결의서
      / 5. 전자세금계산서 / 6. 거래명세서 / 7. 검수확인서
    col A 에 카테고리명(견적서 행), col B 에 "N. 문서명".
  - 업체는 열(column) 헤더. 시트마다 헤더 행/열 범위가 다름:
      Sheet1     : 헤더 row1, cols D~AI
      Sheet1 (2) : 헤더 row1, cols C~S
      결제1/결제2 : 헤더 row3 (결제 금액 추적용 시트)
  - 카테고리명 표기 차이: 체크리스트는 '고효율'/'중동'/'수화후'(오타) 사용.

이 모듈은 (카테고리, 업체, 완료문서목록)을 받아 해당 셀에 'O'를 채우고,
"문서 체크리스트" 시트들(결제 시트 제외)만 기본 대상으로 한다.
"""

import datetime
import difflib
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter


# 문서 종류 -> 블록 내 행 오프셋(견적서=0)
DOC_OFFSETS = {
    "견적서": 0,
    "사업자등록증": 1,
    "통장사본": 2,
    "지출결의서": 3,
    "지결서": 3,
    "전자세금계산서": 4,
    "세금계산서": 4,
    "거래명세서": 5,
    "검수확인서": 6,
}

# 폴더/앱 카테고리 -> 체크리스트 col A 표기
CATEGORY_ALIAS = {
    "고효율": "고효율", "고효율 광원": "고효율",
    "중동": "중동", "IR": "중동", "중동(IR)": "중동",
    "수확후": "수화후", "수화후": "수화후",
    "자동화": "자동화", "북미": "북미", "고온성": "고온성",
    "저온성": "저온성", "근권부": "근권부", "팁스": "팁스",
}

MARK = "O"


def _norm(s):
    return re.sub(r"\s+", "", (s or "")).lower()


def _company_core(s):
    s = re.sub(r"(주식회사|㈜|\(주\)|산업|시스템|테크|이엔지|엔지니어링|인터내셔널)", "", s or "")
    return _norm(s)


# 체크리스트 열 이름과 폴더명이 표기만 다른 경우(영문/한글/음차)를 이어주는 별칭.
# key/값 모두 정규화(_norm) 기준. 양방향으로 매칭에 사용한다.
COMPANY_ALIASES = [
    {"에스씨", "sc"},
    {"에스비에이치이", "sbhe", "sbhe_열교환기", "에스비에이치이_열교환기"},
    {"아이엘패널", "아엘패널", "아이엘패널_판넬", "아엘패널_판넬"},
    {"다임", "다임그린씨에스", "그린씨에스"},
    {"씨엘피", "주식회사씨엘피", "씨엘피_모터"},
]


def _alias_group(token):
    for grp in COMPANY_ALIASES:
        if token in grp:
            return grp
    return {token}


def _find_header_row(ws):
    """업체명 헤더 행 번호를 찾는다(카테고리 첫 블록 위에서 텍스트가 가장 많은 행)."""
    first_cat = None
    for r in range(1, min(ws.max_row, 12) + 1):
        if ws.cell(row=r, column=1).value:
            first_cat = r
            break
    limit = (first_cat or 6) - 1
    best_r, best_n = None, 0
    for r in range(1, max(limit, 1) + 1):
        n = sum(1 for c in range(3, ws.max_column + 1)
                if isinstance(ws.cell(row=r, column=c).value, str)
                and ws.cell(row=r, column=c).value.strip())
        if n > best_n:
            best_r, best_n = r, n
    return best_r


def _company_columns(ws):
    hr = _find_header_row(ws)
    if not hr:
        return {}, None
    cols = {}
    for c in range(3, ws.max_column + 1):
        v = ws.cell(row=hr, column=c).value
        if isinstance(v, str) and v.strip():
            cols[c] = v.strip()
    return cols, hr


def _category_rows(ws):
    """{체크리스트카테고리표기: 견적서행} """
    out = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(row=r, column=1).value
        if isinstance(a, str) and a.strip():
            out[a.strip()] = r
    return out


def _match_company_column(cols, company):
    """company 와 가장 잘 맞는 열 번호를 반환(없으면 None)."""
    target = _norm(company)
    core = _company_core(company)
    # 0) 별칭 그룹(영문/한글/음차 표기 차이) 우선 매칭
    tgroup = _alias_group(target) | _alias_group(core)
    for c, name in cols.items():
        ng = _alias_group(_norm(name)) | _alias_group(_company_core(name))
        if tgroup & ng:
            return c
    # 1) 양방향 부분 문자열
    for c, name in cols.items():
        nn = _norm(name)
        if not nn:
            continue
        if nn in target or target in nn:
            return c
    for c, name in cols.items():
        nn = _company_core(name)
        if nn and core and (nn in core or core in nn) and min(len(nn), len(core)) >= 2:
            return c
    # 2) 유사도(difflib) 0.6 이상 중 최고
    best_c, best_ratio = None, 0.0
    for c, name in cols.items():
        ratio = difflib.SequenceMatcher(None, target, _norm(name)).ratio()
        if ratio > best_ratio:
            best_c, best_ratio = c, ratio
    if best_ratio >= 0.6:
        return best_c
    return None


def update_checklist(src_path, category, company, docs_done, out_path=None,
                     include_payment_sheets=False):
    """체크리스트를 갱신해 새 파일로 저장한다.

    category    : 폴더/앱 카테고리(예: '고효율', '중동')
    company     : 업체명(견적서에서 얻은 정식명도 OK - 축약명과 매칭)
    docs_done   : 완료 문서 목록(예: ['견적서','지출결의서','검수확인서', ...])
    반환: {"out_path", "marked": [ {sheet,cell,doc,company_header} ], "misses": [...]}
    """
    wb = openpyxl.load_workbook(src_path)
    cat_label = CATEGORY_ALIAS.get(category.strip(), category.strip())
    offsets = []
    for d in docs_done:
        key = d.replace(" ", "")
        for name, off in DOC_OFFSETS.items():
            if name == key or name in key:
                offsets.append((d, off))
                break
    # 중복 제거
    seen = set()
    offsets = [(d, o) for d, o in offsets if not (o in seen or seen.add(o))]

    marked, misses = [], []
    any_sheet_hit = False
    for ws in wb.worksheets:
        if not include_payment_sheets and ws.title.startswith("결제"):
            continue
        cols, hr = _company_columns(ws)
        cats = _category_rows(ws)
        if cat_label not in cats:
            continue
        col = _match_company_column(cols, company)
        if col is None:
            continue
        any_sheet_hit = True
        base = cats[cat_label]
        for d, off in offsets:
            r = base + off
            cell = ws.cell(row=r, column=col)
            cell.value = MARK
            marked.append({
                "sheet": ws.title,
                "cell": f"{get_column_letter(col)}{r}",
                "doc": d,
                "company_header": cols.get(col, ""),
            })
    if not any_sheet_hit:
        misses.append(f"체크리스트에서 '{company}' / '{cat_label}' 열·블록을 찾지 못했습니다.")

    if out_path is None:
        stamp = datetime.date.today().strftime("%y%m%d")
        d = os.path.dirname(src_path)
        out_path = os.path.join(d, f"연구비 파일 업로드 체크용_자동갱신_{stamp}.xlsx")
    wb.save(out_path)
    return {"out_path": out_path, "marked": marked, "misses": misses}
