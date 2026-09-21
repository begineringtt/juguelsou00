"""처리 이력 보고서(누적 로그) 작성.

한 업체를 처리할 때마다 한 행을 추가한다. 파일이 이미 있으면 이어 붙이고,
없으면 헤더와 함께 새로 만든다. 문서 제출 추적용이라 xlsx 로 남긴다.
"""

import datetime
import os

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = [
    "처리일시", "과제 카테고리", "업체명", "대상 폴더",
    "생성 문서", "정리한 첨부문서", "견적 총액(원)", "품목 수",
    "견적서 인식방식", "비고",
]

_HEADER_FILL = PatternFill("solid", fgColor="2F5597")
_HEADER_FONT = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="맑은 고딕", size=10)
_ALIGN = Alignment(vertical="center", wrap_text=True)
_THIN = Side(style="thin", color="D0D0D0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_WIDTHS = [18, 12, 20, 30, 22, 28, 14, 8, 14, 24]


def _new_workbook():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "처리이력"
    for i, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = _HEADER_FILL
        c.font = _HEADER_FONT
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = _WIDTHS[i - 1]
    ws.freeze_panes = "A2"
    return wb


def append_entry(report_path, entry):
    """entry(dict)를 보고서에 한 행 추가한다. entry 키는 아래 순서에 대응.

    entry = {
      category, company, folder, generated (list|str), attachments (list|str),
      total_amount (int|None), item_count (int|None), source (str), note (str)
    }
    """
    if os.path.isfile(report_path):
        wb = openpyxl.load_workbook(report_path)
        ws = wb["처리이력"] if "처리이력" in wb.sheetnames else wb.active
    else:
        wb = _new_workbook()
        ws = wb.active

    def _join(v):
        if isinstance(v, (list, tuple)):
            return ", ".join(str(x) for x in v)
        return "" if v is None else str(v)

    row = [
        datetime.datetime.now().strftime("%Y-%m-%d %H:%M"),
        _join(entry.get("category")),
        _join(entry.get("company")),
        _join(entry.get("folder")),
        _join(entry.get("generated")),
        _join(entry.get("attachments")),
        entry.get("total_amount"),
        entry.get("item_count"),
        _join(entry.get("source")),
        _join(entry.get("note")),
    ]
    r = ws.max_row + 1
    for i, val in enumerate(row, start=1):
        c = ws.cell(row=r, column=i, value=val)
        c.font = _BODY_FONT
        c.alignment = _ALIGN
        c.border = _BORDER
        if i == 7 and isinstance(val, (int, float)):
            c.number_format = "#,##0"
    wb.save(report_path)
    return report_path
