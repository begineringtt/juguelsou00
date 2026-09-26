"""연구비 소진 계획 대비 지출결의서 진행 현황 체크.

"연구비 소진 계획.xlsx"(과제x업체 계획 금액 매트릭스)와 실제 setting_03 폴더 현황을
대조해 (1) 계획은 있는데 아직 지출결의서를 안 만든 (과제,업체)와 (2) 이미 만들었지만
실제 처리 금액이 계획과 다른 건을 찾는다. "01. 지출결의서_전체과제_통합(양식).xlsx"는
과제 목록(시트 이름) 참조용으로만 쓴다 - 문서 내부 구조까지는 비교하지 않는다.
"""

import os

import openpyxl

import checklist_sync
import checklist_updater as CU
import folder_router

# 연구비 소진 계획.xlsx 의 과제 표기가 실제 setting_03 폴더명과 전혀 다르게 지어진
# 경우(folder_router.CATEGORY_TO_FOLDER 의 부분 문자열 별칭으로도 못 잡는 경우)만
# 여기서 별도로 잡아준다.
PROJECT_ALIASES = {
    "AI(1년)": "AI(1y)그린CS",
    "AI(1.5년)": "AI(1.5y)메타파머스",
    # "01. 지출결의서_전체과제_통합(양식).xlsx"의 시트 이름은 또 다른 표기를 쓴다.
    "AICS1년": "AI(1y)그린CS",
    "AI 메타1.5년": "AI(1.5y)메타파머스",
}


def resolve_folder(project_label):
    if project_label in PROJECT_ALIASES:
        return PROJECT_ALIASES[project_label]
    return folder_router.folder_for_category(project_label)


PLAN_SHEET_NAME = "결제금액 계획"
PLAN_COMPANY_HEADER_ROW = 1
PLAN_FIRST_PROJECT_ROW = 3
PLAN_COMPANY_START_COL = 4  # D열

# 과제 열(B열)에 나타나지만 실제 과제가 아닌 요약 행 라벨.
PLAN_NON_PROJECT_LABELS = {"소계", "오차금액"}


def _plan_company_columns(ws):
    """D열부터 시작해, 헤더(1행)에 연속으로 빈 칸이 2개 나오면 멈춘다 - 매트릭스
    오른쪽 끝에 붙은 별개의 메모/범례 칸(예: "통장사본,사업자ok")까지 회사로
    잘못 집어내지 않기 위함."""
    cols = []
    misses = 0
    col = PLAN_COMPANY_START_COL
    while col <= ws.max_column and misses < 2:
        value = ws.cell(row=PLAN_COMPANY_HEADER_ROW, column=col).value
        if isinstance(value, str) and value.strip():
            cols.append(col)
            misses = 0
        else:
            misses += 1
        col += 1
    return cols


DEFAULT_PLAN_FILENAME = "연구비 소진 계획.xlsx"
DEFAULT_MASTER_FILENAME = "01. 지출결의서_전체과제_통합(양식).xlsx"
DEFAULT_REPORT_RELPATH = os.path.join("체크리스트", "처리이력_보고서.xlsx")


def check_progress(setting03_root, plan_path=None, master_path=None, report_path=None):
    """"연구비 소진 계획.xlsx" 대비 아직 안 만든 (과제,업체)와 금액이 다른 건,
    "01. 지출결의서_전체과제_통합(양식).xlsx" 기준으로 문서가 하나도 없는 과제를
    찾는다. 세 경로 모두 생략하면 setting03_root 밑의 기본 위치를 쓴다.

    반환: {missing, missing_count, mismatches, projects_without_any_report}
    """
    plan_path = plan_path or os.path.join(setting03_root, DEFAULT_PLAN_FILENAME)
    master_path = master_path or os.path.join(setting03_root, DEFAULT_MASTER_FILENAME)
    report_path = report_path or os.path.join(setting03_root, DEFAULT_REPORT_RELPATH)

    plan_rows = load_plan(plan_path)
    actual_totals = load_actual_totals(report_path)
    master_folders = sorted({resolve_folder(p) for p in load_master_projects(master_path)})

    all_folders = sorted({row["folder"] for row in plan_rows} | set(master_folders))
    scanned_by_folder = {}
    for entry in checklist_sync.scan_status(setting03_root, categories=all_folders):
        scanned_by_folder.setdefault(entry["category_folder"], []).append(entry)

    totals_by_folder = {}
    for (folder, company), amount in actual_totals.items():
        totals_by_folder.setdefault(folder, []).append((company, amount))

    missing, mismatches = [], []
    for row in plan_rows:
        folder = row["folder"]
        entries = scanned_by_folder.get(folder, [])
        cols = {i: e["company"] for i, e in enumerate(entries)}
        idx = CU._match_company_column(cols, row["company"])
        has_report = idx is not None and "지출결의서" in entries[idx]["docs"]
        if not has_report:
            missing.append({
                "project": row["project"], "folder": folder,
                "company": row["company"], "planned_amount": row["planned_amount"],
            })
            continue

        amount_entries = totals_by_folder.get(folder, [])
        amount_cols = {i: c for i, (c, _a) in enumerate(amount_entries)}
        amount_idx = CU._match_company_column(amount_cols, row["company"])
        if amount_idx is None:
            continue
        actual_amount = amount_entries[amount_idx][1]
        diff = actual_amount - row["planned_amount"]
        if diff != 0:
            mismatches.append({
                "project": row["project"], "folder": folder, "company": row["company"],
                "planned_amount": row["planned_amount"], "actual_amount": actual_amount,
                "diff": diff,
            })

    projects_without_any_report = [
        folder for folder in master_folders
        if not any("지출결의서" in e["docs"] for e in scanned_by_folder.get(folder, []))
    ]

    return {
        "missing": missing,
        "missing_count": len(missing),
        "mismatches": mismatches,
        "projects_without_any_report": projects_without_any_report,
    }


def scaffold_folders(setting03_root, plan_path=None):
    """"연구비 소진 계획.xlsx"에 있는 (과제,업체) 조합마다 업체 하위 폴더를
    미리 만들어 둔다. 이미 비슷한 이름의 폴더가 있으면(날짜/차수 접두 등이
    붙은 기존 폴더 포함) 새로 만들지 않고 건너뛴다.

    반환: {"created": [{folder,company,path}], "already_existed": [{folder,company,matched_folder}]}
    """
    plan_path = plan_path or os.path.join(setting03_root, DEFAULT_PLAN_FILENAME)
    plan_rows = load_plan(plan_path)

    seen = set()
    created, already_existed = [], []
    for row in plan_rows:
        folder, company = row["folder"], row["company"]
        if (folder, company) in seen:
            continue
        seen.add((folder, company))

        folder_root = os.path.join(setting03_root, folder)
        existing = [n for n in os.listdir(folder_root) if os.path.isdir(os.path.join(folder_root, n))] \
            if os.path.isdir(folder_root) else []
        cols = dict(enumerate(existing))
        idx = CU._match_company_column(cols, company)
        if idx is not None:
            already_existed.append({"folder": folder, "company": company, "matched_folder": existing[idx]})
            continue

        path = os.path.join(folder_root, company)
        os.makedirs(path, exist_ok=True)
        created.append({"folder": folder, "company": company, "path": path})

    return {"created": created, "already_existed": already_existed}


def load_master_projects(master_path):
    """"01. 지출결의서_전체과제_통합(양식).xlsx"의 시트 이름(과제 목록)을 그대로
    반환한다 - 이 문서의 내부 구조까지는 비교하지 않고, "이 과제들이 있어야
    한다"는 목록 참조로만 쓴다."""
    wb = openpyxl.load_workbook(master_path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def load_actual_totals(report_path):
    """"처리이력_보고서.xlsx"(report_writer.append_entry 가 매번 남기는 로그)를
    읽어 (과제 폴더, 업체명)별 실제 처리 총액 합계를 만든다.

    반환: {(folder, company): summed_total_amount} - 로그가 없으면 {}.
    """
    if not os.path.isfile(report_path):
        return {}
    wb = openpyxl.load_workbook(report_path, data_only=True)
    ws = wb["처리이력"] if "처리이력" in wb.sheetnames else wb.active

    totals = {}
    for r in range(2, ws.max_row + 1):
        category = ws.cell(row=r, column=2).value
        company = ws.cell(row=r, column=3).value
        amount = ws.cell(row=r, column=7).value
        if not isinstance(category, str) or not isinstance(company, str):
            continue
        if not isinstance(amount, (int, float)):
            continue
        key = (resolve_folder(category.strip()), company.strip())
        totals[key] = totals.get(key, 0) + amount
    return totals


def load_plan(plan_path):
    """"연구비 소진 계획.xlsx"를 읽어 계획 금액이 있는 (과제,업체) 목록을 만든다.

    반환: [ {project, folder, company, planned_amount}, ... ] (금액이 없거나
    0인 칸은 제외 - "계획된 적 없음"으로 본다.)
    """
    wb = openpyxl.load_workbook(plan_path, data_only=True)
    ws = wb[PLAN_SHEET_NAME] if PLAN_SHEET_NAME in wb.sheetnames else wb.active
    company_cols = _plan_company_columns(ws)
    companies = {col: ws.cell(row=PLAN_COMPANY_HEADER_ROW, column=col).value.strip()
                 for col in company_cols}

    rows = []
    for r in range(PLAN_FIRST_PROJECT_ROW, ws.max_row + 1):
        project = ws.cell(row=r, column=2).value
        if not isinstance(project, str) or not project.strip():
            continue
        project = project.strip()
        if project in PLAN_NON_PROJECT_LABELS:
            continue
        for col, company in companies.items():
            amount = ws.cell(row=r, column=col).value
            if isinstance(amount, (int, float)) and amount > 0:
                rows.append({
                    "project": project,
                    "folder": resolve_folder(project),
                    "company": company,
                    "planned_amount": amount,
                })
    return rows
